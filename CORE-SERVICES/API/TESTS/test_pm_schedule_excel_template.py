"""AI5R-PHASE4E4, Section D/O -- the downloadable PM Schedule import
template: exact required columns (no Schedule Code/Status/Trigger Type/
Procedure/Checklist Template/audit fields), a frozen header row, an
Instructions sheet that explains accepted values without ever mentioning
WCH, and the formula-leading-character escape helper (Section M)."""

import io
import sys
from pathlib import Path

from openpyxl import load_workbook

CORE_SERVICES_PATH = Path(__file__).resolve().parents[2]
if str(CORE_SERVICES_PATH) not in sys.path:
    sys.path.insert(0, str(CORE_SERVICES_PATH))

from API.pm_schedule_excel_template import (  # noqa: E402
    TEMPLATE_COLUMNS,
    build_pm_schedule_import_template,
    escape_formula_leading,
)


def _load(content: bytes):
    return load_workbook(io.BytesIO(content))


def test_template_has_exactly_the_required_columns_in_order():
    workbook = _load(build_pm_schedule_import_template())
    sheet = workbook["PM Schedule Import"]
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header == list(TEMPLATE_COLUMNS)


def test_template_never_includes_forbidden_system_managed_columns():
    workbook = _load(build_pm_schedule_import_template())
    sheet = workbook["PM Schedule Import"]
    header = " ".join(str(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    for forbidden in ("Schedule Code", "Status", "Trigger Type", "Procedure", "Checklist Template"):
        assert forbidden not in header


def test_data_sheet_has_no_data_rows_only_the_header():
    workbook = _load(build_pm_schedule_import_template())
    sheet = workbook["PM Schedule Import"]
    assert sheet.max_row == 1


def test_header_row_is_frozen():
    workbook = _load(build_pm_schedule_import_template())
    sheet = workbook["PM Schedule Import"]
    assert sheet.freeze_panes == "A2"


def test_instructions_sheet_exists_and_explains_required_fields():
    workbook = _load(build_pm_schedule_import_template())
    assert "Instructions" in workbook.sheetnames
    text = "\n".join(
        " ".join(str(c.value) for c in row if c.value is not None)
        for row in workbook["Instructions"].iter_rows()
    )
    assert "Pump Tag" in text
    assert "Frequency" in text
    assert "Start Date" in text
    assert "DAILY" in text
    assert "RUNTIME_BASED" in text


def test_instructions_explain_cooler_vs_cooling_water_cooler_and_reservoir():
    workbook = _load(build_pm_schedule_import_template())
    text = "\n".join(
        " ".join(str(c.value) for c in row if c.value is not None)
        for row in workbook["Instructions"].iter_rows()
    )
    assert "Cooling Water Cooler" in text
    assert "different piece of equipment" in text or "never the same row" in text
    assert "Reservoir has no DE/NDE" in text or "General only" in text


def test_instructions_never_mention_wch():
    workbook = _load(build_pm_schedule_import_template())
    text = "\n".join(
        " ".join(str(c.value) for c in row if c.value is not None)
        for row in workbook["Instructions"].iter_rows()
    )
    assert "WCH" not in text.upper().replace(" ", "")


def test_escape_formula_leading_neutralizes_every_dangerous_prefix():
    for prefix in ("=", "+", "-", "@"):
        assert escape_formula_leading(f"{prefix}cmd").startswith("'")
    assert escape_formula_leading("normal text") == "normal text"
    assert escape_formula_leading("") == ""
