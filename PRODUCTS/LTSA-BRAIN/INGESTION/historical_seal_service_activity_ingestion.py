"""MWO-LTSA-HISTORICAL-SEAL-SERVICE-ACTIVITY-001 -- ingestion engine for
historical mechanical seal service activity records (2024-2025).

MANDATORY RULES:
1. HISTORY ONLY: Historical service activities are chronological lifecycle evidence;
   they MUST NOT establish or overwrite Current Installation snapshots.
2. CANONICAL DATE: Finish Date is the authoritative event_date. Start Date or
   Failure Date may remain as raw metadata but never replace missing Finish Date.
3. CONSERVATIVE ASSET MATCHING: Only EXACT and NORMALIZED_EXACT tags are linked
   to pump_tag_number (FK to public.ltsa_pumps). AMBIGUOUS (missing suffix)
   and NO_MATCH tags remain unlinked (pump_tag_number IS NULL) for human review.
4. IDEMPOTENT: Enforced by UNIQUE (source_reference) and ON CONFLICT DO NOTHING.
   Re-running import causes zero duplicate entries.
"""

from __future__ import annotations

import argparse
import datetime
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

import openpyxl

MONTH_MAP = {
    "JAN": 1, "JANUARI": 1, "JANUARY": 1,
    "FEB": 2, "FEBRUARI": 2, "FEBRUARY": 2,
    "MAR": 3, "MARET": 3, "MARCH": 3,
    "APR": 4, "APRIL": 4,
    "MEI": 5, "MAY": 5,
    "JUN": 6, "JUNI": 6, "JUNE": 6,
    "JUL": 7, "JULI": 7, "JULY": 7,
    "AGU": 8, "AGUSTUS": 8, "AUG": 8, "AUGUST": 8,
    "SEP": 9, "SEPTEMBER": 9,
    "OKT": 10, "OKTOBER": 10, "OCT": 10, "OCTOBER": 10,
    "NOV": 11, "NOVEMBER": 11,
    "DES": 12, "DESEMBER": 12, "DEC": 12, "DECEMBER": 12,
}


def parse_date_val(val: Any) -> tuple[datetime.date | None, str]:
    """Parse date from cell value.
    Returns (parsed_date, date_status).
    If val is None, blank, or placeholder ('-'), returns (None, 'MISSING')."""
    if val is None:
        return None, "MISSING"
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.date() if isinstance(val, datetime.datetime) else val, "VALID"
    text = str(val).strip()
    if text in ("", "-", "N/A", "None", "NULL"):
        return None, "MISSING"

    m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$", text)
    if m:
        d, mth, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return datetime.date(y, mth, d), "VALID"
        except ValueError:
            return None, f"INVALID_{text}"

    m2 = re.match(r"^(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})$", text)
    if m2:
        d = int(m2.group(1))
        mon_str = m2.group(2).upper()
        y = int(m2.group(3))
        if mon_str in MONTH_MAP:
            mth = MONTH_MAP[mon_str]
            try:
                return datetime.date(y, mth, d), "VALID"
            except ValueError:
                return None, f"INVALID_{text}"
    return None, f"INVALID_{text}"


def classify_tag(raw_tag: Any, canonical_tags: set[str]) -> tuple[str, str | None, str]:
    """Classify raw tag against canonical pump registry tags.
    Outcomes:
      EXACT: exact string match
      NORMALIZED_EXACT: case, whitespace, or (NDE)/(DE)/SPARE stripped match
      AMBIGUOUS: near-miss missing suffix matching canonical candidate(s) (e.g. 211-P-14 -> 211-P-14A/B)
      NO_MATCH: no candidate in canonical registry
      MISSING_TAG: tag is null or placeholder
    Conservative rule: pump_tag_number is ONLY populated for EXACT and NORMALIZED_EXACT.
    """
    if raw_tag is None:
        return "MISSING_TAG", None, "Tag is NULL"
    tag = str(raw_tag).strip()
    if not tag or tag in ("-", "N/A"):
        return "MISSING_TAG", None, f"Tag is blank or placeholder: {raw_tag}"

    if tag in canonical_tags:
        return "EXACT", tag, "Exact match"

    tag_upper = tag.upper()
    if tag_upper in canonical_tags:
        return "NORMALIZED_EXACT", tag_upper, "Case normalized"

    collapsed = re.sub(r"\s+", "", tag_upper)
    if collapsed in canonical_tags:
        return "NORMALIZED_EXACT", collapsed, "Whitespace collapsed"

    cleaned = re.sub(r"\s*\((?:NDE|DE)\)\s*$", "", tag_upper).strip()
    cleaned = re.sub(r"\s+SPARE\s*$", "", cleaned).strip()
    if cleaned in canonical_tags:
        return "NORMALIZED_EXACT", cleaned, f"Stripped position/spare qualifier: {tag} -> {cleaned}"

    base_tag = re.sub(r"[A-Z]+$", "", collapsed)
    matching_canonical = [
        ct for ct in canonical_tags
        if re.sub(r"[A-Z]+$", "", re.sub(r"\s+", "", ct)) == base_tag
    ]
    if len(matching_canonical) >= 1:
        return "AMBIGUOUS", None, f"Near-miss missing suffix matching canonical candidate(s): {matching_canonical}"

    return "NO_MATCH", None, f"No match in canonical pump registry: {tag}"


def extract_2024_records(wb_path: Path, canonical_tags: set[str]) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    records = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        h_row = 7
        headers = [ws.cell(h_row, c).value for c in range(1, ws.max_column + 1)]
        has_double_no = (len(headers) > 1 and headers[0] == "No" and headers[1] == "No")

        for r in range(h_row + 1, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            tag = vals[3] if has_double_no else vals[2]
            desc = vals[4] if has_double_no else vals[3]
            seal = vals[6] if has_double_no else vals[5]
            size = vals[7] if has_double_no else vals[6]
            start_d = vals[13] if has_double_no else vals[12]
            finish_d = vals[14] if has_double_no else vals[13]
            job_no = vals[2] if has_double_no else vals[1]
            team = vals[5] if has_double_no else vals[4]
            sleeve = vals[8] if has_double_no else vals[7]
            gland = vals[9] if has_double_no else vals[8]
            qty = vals[10] if has_double_no else vals[9]
            dwg = vals[11] if has_double_no else vals[10]
            end_user = vals[12] if has_double_no else vals[11]
            sp_no = vals[15] if has_double_no else vals[14]
            status = vals[16] if has_double_no else vals[15]
            remarks = vals[17] if has_double_no else vals[16]

            if not any(x is not None and str(x).strip() not in ("", "-") for x in (tag, desc, finish_d)):
                continue

            outcome, matched, reason = classify_tag(tag, canonical_tags)
            parsed_finish, finish_status = parse_date_val(finish_d)
            parsed_start, _ = parse_date_val(start_d)

            # Canonical event date is ALWAYS finish date (Never substitute start date)
            event_date = parsed_finish

            records.append({
                "source_reference": f"SERVICE_ACTIVITY:2024:{sheet_name}:R{r}:{tag}",
                "source_type": "SERVICE_ACTIVITY",
                "source_year": 2024,
                "source_filename": wb_path.name,
                "source_worksheet": sheet_name,
                "source_row": r,
                "source_job_document_number": str(job_no).strip() if job_no is not None else None,
                "sp_no": str(sp_no).strip() if sp_no is not None else None,
                "raw_tag": str(tag).strip() if tag is not None else "",
                "pump_tag_number": matched,
                "tag_match_outcome": outcome,
                "raw_job_description": str(desc).strip() if desc is not None else None,
                "event_type": "INSTALLATION",
                "failure_attribution": "UNKNOWN",
                "seal_type": str(seal).strip() if seal is not None else None,
                "seal_size": str(size).strip() if size is not None else None,
                "drawing_number": str(dwg).strip() if dwg is not None else None,
                "quantity": int(qty) if qty is not None and str(qty).strip().isdigit() else 1,
                "shaft_sleeve_condition": str(sleeve).strip() if sleeve is not None else None,
                "gland_plate_condition": str(gland).strip() if gland is not None else None,
                "team_service": str(team).strip() if team is not None else None,
                "end_user": str(end_user).strip() if end_user is not None else None,
                "location": None,
                "ltsa_area": None,
                "unit_area": None,
                "api_plan": None,
                "pump_type": None,
                "process_fluid": None,
                "source_start_date": parsed_start,
                "source_failure_date": None,
                "finish_date": parsed_finish,
                "event_date": event_date,
                "date_status": finish_status,
                "status": str(status).strip() if status is not None else None,
                "remarks": str(remarks).strip() if remarks is not None else None,
            })
    return records


def extract_2025_records(wb_path: Path, canonical_tags: set[str]) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ws = wb["INSTALLATION REPORT 2025"]
    records = []
    for r in range(2, 63):
        vals = [ws.cell(r, c).value for c in range(1, 25)]
        no, month, desc, ltsa_area, unit_area, doc_no, tag, seal, size, api_plan, pump_type, process_fluid, dwg, basic_seal, sleeve, gland, location, team, fail_d, finish_d, sp_no, packing_list, status, remarks = vals

        outcome, matched, reason = classify_tag(tag, canonical_tags)
        parsed_finish, finish_status = parse_date_val(finish_d)
        parsed_fail, _ = parse_date_val(fail_d)

        # Canonical event date is ALWAYS finish date (Never substitute failure date)
        event_date = parsed_finish

        records.append({
            "source_reference": f"SERVICE_ACTIVITY:2025:INSTALLATION REPORT 2025:R{r}:{tag}",
            "source_type": "SERVICE_ACTIVITY",
            "source_year": 2025,
            "source_filename": wb_path.name,
            "source_worksheet": "INSTALLATION REPORT 2025",
            "source_row": r,
            "source_job_document_number": str(doc_no).strip() if doc_no is not None else None,
            "sp_no": str(sp_no).strip() if sp_no is not None else None,
            "raw_tag": str(tag).strip() if tag is not None else "",
            "pump_tag_number": matched,
            "tag_match_outcome": outcome,
            "raw_job_description": str(desc).strip() if desc is not None else None,
            "event_type": "INSTALLATION",
            "failure_attribution": "UNKNOWN",
            "seal_type": str(seal).strip() if seal is not None else None,
            "seal_size": str(size).strip() if size is not None else None,
            "drawing_number": str(dwg).strip() if dwg is not None else None,
            "quantity": 1,
            "shaft_sleeve_condition": str(sleeve).strip() if sleeve is not None else None,
            "gland_plate_condition": str(gland).strip() if gland is not None else None,
            "team_service": str(team).strip() if team is not None else None,
            "end_user": None,
            "location": str(location).strip() if location is not None else None,
            "ltsa_area": str(ltsa_area).strip() if ltsa_area is not None else None,
            "unit_area": str(unit_area).strip() if unit_area is not None else None,
            "api_plan": str(api_plan).strip() if api_plan is not None else None,
            "pump_type": str(pump_type).strip() if pump_type is not None else None,
            "process_fluid": str(process_fluid).strip() if process_fluid is not None else None,
            "source_start_date": None,
            "source_failure_date": parsed_fail,
            "finish_date": parsed_finish,
            "event_date": event_date,
            "date_status": finish_status,
            "status": str(status).strip() if status is not None else None,
            "remarks": str(remarks).strip() if remarks is not None else None,
        })
    return records


def extract_all_records(wb_2024_path: Path, wb_2025_path: Path, canonical_tags: set[str]) -> list[dict[str, Any]]:
    records_2024 = extract_2024_records(wb_2024_path, canonical_tags)
    records_2025 = extract_2025_records(wb_2025_path, canonical_tags)
    return records_2024 + records_2025


def _sql_val(val: Any) -> str:
    if val is None:
        return "NULL"
    if isinstance(val, (int, float)):
        return str(val)
    if isinstance(val, datetime.date):
        return f"'{val.isoformat()}'"
    return "'" + str(val).replace("'", "''") + "'"


def build_insert_sql(record: dict[str, Any]) -> str:
    columns = [
        "source_reference", "source_type", "source_year", "source_filename",
        "source_worksheet", "source_row", "source_job_document_number", "sp_no",
        "raw_tag", "pump_tag_number", "tag_match_outcome", "raw_job_description",
        "event_type", "failure_attribution", "seal_type", "seal_size",
        "drawing_number", "quantity", "shaft_sleeve_condition", "gland_plate_condition",
        "team_service", "end_user", "location", "ltsa_area", "unit_area", "api_plan",
        "pump_type", "process_fluid", "source_start_date", "source_failure_date",
        "finish_date", "event_date", "date_status", "status", "remarks"
    ]
    vals = [_sql_val(record.get(col)) for col in columns]
    return (
        f"INSERT INTO public.historical_seal_service_activity ({', '.join(columns)}) "
        f"VALUES ({', '.join(vals)}) "
        "ON CONFLICT (source_reference) DO NOTHING;"
    )


def fetch_canonical_pump_tags(runner: Any = None) -> set[str]:
    """Fetch canonical pump tags from database."""
    if runner is not None and hasattr(runner, "query_scalar"):
        raw = runner.query_scalar("SELECT json_agg(tag_number)::text FROM ltsa_pumps;")
        if raw:
            return set(json.loads(raw))
    res = subprocess.run(
        ["docker", "exec", "-i", "ai5r-runtime-postgres-1", "psql", "-U", "ai5r", "-d", "ltsa_brain", "-t", "-A", "-c",
         "SELECT tag_number FROM ltsa_pumps ORDER BY tag_number;"],
        capture_output=True, text=True, check=True
    )
    return set(line.strip() for line in res.stdout.strip().split("\n") if line.strip())


def run_sql_script(sql: str, runner: Any = None) -> None:
    if runner is not None and hasattr(runner, "execute_script"):
        runner.execute_script(sql)
        return
    subprocess.run(
        ["docker", "exec", "-i", "ai5r-runtime-postgres-1", "psql", "-U", "ai5r", "-d", "ltsa_brain", "-v", "ON_ERROR_STOP=1"],
        input=sql, text=True, check=True, capture_output=True
    )


def execute_ingestion(
    records: list[dict[str, Any]],
    runner: Any = None,
    dry_run: bool = False,
) -> dict[str, Any]:
    total_source_rows = len(records)
    finish_date_valid = sum(1 for r in records if r["date_status"] == "VALID")
    finish_date_missing = sum(1 for r in records if r["date_status"] == "MISSING")
    asset_exact = sum(1 for r in records if r["tag_match_outcome"] == "EXACT")
    asset_normalized_exact = sum(1 for r in records if r["tag_match_outcome"] == "NORMALIZED_EXACT")
    asset_ambiguous = sum(1 for r in records if r["tag_match_outcome"] == "AMBIGUOUS")
    asset_no_match = sum(1 for r in records if r["tag_match_outcome"] == "NO_MATCH")

    stats = {
        "SOURCE_ROWS": total_source_rows,
        "FINISH_DATE_VALID": finish_date_valid,
        "FINISH_DATE_MISSING": finish_date_missing,
        "ASSET_EXACT": asset_exact,
        "ASSET_NORMALIZED_EXACT": asset_normalized_exact,
        "ASSET_AMBIGUOUS": asset_ambiguous,
        "ASSET_NO_MATCH": asset_no_match,
        "INSERTED_ROWS": 0,
        "SKIPPED_ROWS": 0,
        "DRY_RUN": dry_run,
    }

    if dry_run:
        return stats

    # Get count before
    count_sql = "SELECT COUNT(*) FROM public.historical_seal_service_activity;"
    count_before = 0
    if runner is not None and hasattr(runner, "query_scalar"):
        count_before = int(runner.query_scalar(count_sql) or 0)
    else:
        res = subprocess.run(
            ["docker", "exec", "-i", "ai5r-runtime-postgres-1", "psql", "-U", "ai5r", "-d", "ltsa_brain", "-t", "-A", "-c", count_sql],
            capture_output=True, text=True, check=True
        )
        count_before = int(res.stdout.strip() or 0)

    sql_statements = ["BEGIN;"]
    for r in records:
        sql_statements.append(build_insert_sql(r))
    sql_statements.append("COMMIT;")
    full_script = "\n".join(sql_statements)

    run_sql_script(full_script, runner)

    # Get count after
    count_after = 0
    if runner is not None and hasattr(runner, "query_scalar"):
        count_after = int(runner.query_scalar(count_sql) or 0)
    else:
        res = subprocess.run(
            ["docker", "exec", "-i", "ai5r-runtime-postgres-1", "psql", "-U", "ai5r", "-d", "ltsa_brain", "-t", "-A", "-c", count_sql],
            capture_output=True, text=True, check=True
        )
        count_after = int(res.stdout.strip() or 0)

    inserted = count_after - count_before
    stats["INSERTED_ROWS"] = inserted
    stats["SKIPPED_ROWS"] = total_source_rows - inserted
    stats["TOTAL_ROWS_AFTER"] = count_after
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description="Historical Seal Service Activity Ingestion (2024-2025)")
    parser.add_argument(
        "--wb-2024",
        type=Path,
        default=Path(r"C:\Users\USER\Downloads\Servis Activity 2024 (1).xlsx"),
        help="Path to 2024 Service Activity workbook",
    )
    parser.add_argument(
        "--wb-2025",
        type=Path,
        default=Path(r"C:\Users\USER\Downloads\SERVICE ACTIFITY JAN - DES 2025 (1).xlsx"),
        help="Path to 2025 Service Activity workbook",
    )
    parser.add_argument("--apply", action="store_true", help="Execute actual database insertion")
    parser.add_argument("--dry-run", action="store_true", help="Dry run without modifying database")

    args = parser.parse_args()
    dry_run = not args.apply

    canonical_tags = fetch_canonical_pump_tags()
    records = extract_all_records(args.wb_2024, args.wb_2025, canonical_tags)

    stats = execute_ingestion(records, dry_run=dry_run)
    print(json.dumps(stats, indent=2))


if __name__ == "__main__":
    main()
